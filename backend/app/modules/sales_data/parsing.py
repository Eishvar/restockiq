"""Turn an uploaded CSV or Excel file into plain rows of text. No database and no AI here.

Both formats are first read into a grid of text cells. The header row is then found by its column names, so a title
line or blank rows above the table (common in till exports) do no harm. Nothing is guessed: a cell that cannot be
understood is passed on as text and rejected, with a reason, by the next step.
"""
import csv
import io
import re
from datetime import date, datetime

from fastapi import HTTPException

MAX_BYTES = 8_000_000
MAX_ROWS = 60_000

# Header names are compared in lower case with everything but letters and digits removed.
ALIASES = {
    "date": {"date", "salesdate", "businessdate", "day", "tarikh"},
    "code": {"itemcode", "itemid", "code", "sku", "plu"},
    "name": {"itemname", "item", "name", "product", "dish", "menuitem", "description"},
    "qty": {"qtysold", "qty", "quantity", "quantitysold", "sold", "units"},
    "net": {"netsalesrm", "netsales", "total", "amount", "sales", "salesrm", "revenue"},
}

KINDS = {".csv": "csv", ".txt": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".pdf": "pdf"}

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%d-%b-%Y", "%d/%m/%y", "%b %d, %Y")


def kind_of(filename: str) -> str:
    name = filename.lower()
    if name.endswith(".xls"):
        raise HTTPException(422, "Old .xls files are not supported. Save the file as .xlsx or CSV and upload that.")
    for ext, kind in KINDS.items():
        if name.endswith(ext):
            return kind
    raise HTTPException(422, "Upload a CSV, Excel (.xlsx) or PDF file.")


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


def parse_date(text: str) -> date:
    text = " ".join(str(text).split())
    text = re.sub(r"[ T]00:00:00(\.0+)?$", "", text)  # an Excel date cell read as text
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("date not understood (use YYYY-MM-DD or DD/MM/YYYY)")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _csv_grid(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(422, "That does not look like a CSV file")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _xlsx_grid(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(422, "That does not look like an Excel (.xlsx) file") from None
    best: list[list[str]] = []
    for sheet in book.worksheets:  # the first sheet that has a recognisable table wins
        grid = [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
        if _find_header(grid) is not None:
            return grid
        best = best or grid
    return best


def _find_header(grid: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    """The first of the top 30 rows that names a date, a quantity and a dish (by code or name)."""
    for i, row in enumerate(grid[:30]):
        cols: dict[str, int] = {}
        for j, header in enumerate(row):
            for key, names in ALIASES.items():
                if _norm(header) in names and key not in cols:
                    cols[key] = j
        if "date" in cols and "qty" in cols and ({"code", "name"} & cols.keys()):
            return i, cols
    return None


def read_rows(filename: str, data: bytes) -> tuple[str, list[dict]]:
    """Returns (kind, rows). Each row: line (the line in the file), date, code, dish, qty, net (all text)."""
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "That file is too big. Upload a shorter date range.")
    kind = kind_of(filename)
    if kind == "pdf":
        raise ValueError("PDF files are read by the AI step, not here")
    grid = _csv_grid(data) if kind == "csv" else _xlsx_grid(data)
    found = _find_header(grid)
    if found is None:
        raise HTTPException(
            422,
            "The file needs these columns: Date, Item Code or Item Name, and Qty Sold. "
            f"Found: {', '.join(c for c in (grid[0] if grid else []) if c) or 'none'}",
        )
    header_at, cols = found
    rows = []
    for line, row in enumerate(grid[header_at + 1 :], start=header_at + 2):
        if not any(c.strip() for c in row):
            continue

        def get(key: str) -> str:
            j = cols.get(key)
            return row[j].strip() if j is not None and j < len(row) else ""

        rows.append({"line": line, "date": get("date"), "code": get("code"), "dish": get("name"),
                     "qty": get("qty"), "net": get("net")})
        if len(rows) > MAX_ROWS:
            raise HTTPException(413, "That file has too many rows. Upload a shorter date range.")
    if not rows:
        raise HTTPException(422, "The file has a header but no sales rows.")
    return kind, rows
